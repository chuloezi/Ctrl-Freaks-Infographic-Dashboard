"""Ctrl Freaks | reproducible Rwanda CRVS analytics.

Run: python analytics/analyze_crvs.py
Input: three original NISR final-tables XLSX workbooks in data/raw.
Output: transparent CSV source rows, analysis tables, chart, and validation report.

Uses openpyxl to read XLSX, stdlib CSV and matplotlib only for the optional
plot. Never silently imputes missing values.
"""
from __future__ import annotations
import csv, json, re
from collections import Counter
from pathlib import Path

ROOT=Path(__file__).resolve().parents[1]
RAW=ROOT/'data'/'raw'
OUT=ROOT/'data'/'processed'
OUT.mkdir(parents=True,exist_ok=True)

SOURCE_FILES={
  2023:'Final Tables_CRVS_2023.xlsx',
  2024:'Final Tables_CRVS_2024_N.xlsx',
  2025:'Final Tables_CRVS_2025_0.xlsx',
}

def write_csv(path, columns, records):
    with path.open('w',newline='',encoding='utf-8') as f:
        writer=csv.DictWriter(f,fieldnames=columns)
        writer.writeheader()
        writer.writerows(records)

def clean(v):
    if v is None:return ''
    if isinstance(v,float):return f'{v:.12g}'
    return str(v).strip().replace('\n',' | ')

def read_workbooks():
    from openpyxl import load_workbook
    records=[]; figures={}; sheet_index=[]
    for publication_year, name in SOURCE_FILES.items():
        path=RAW/name
        if not path.exists():raise FileNotFoundError('Missing original NISR source: '+str(path))
        wb=load_workbook(path,data_only=True)
        for sheet in wb.worksheets:
            first_row=sheet.min_row
            address=sheet.dimensions
            sheet_rows=[[clean(v) for v in row] for row in sheet.iter_rows(values_only=True,min_row=sheet.min_row,max_row=sheet.max_row,min_col=sheet.min_column,max_col=sheet.max_column)]
            if publication_year==2025:
                figures[sheet.title.strip()]=sheet_rows
            count=0
            for offset, cells in enumerate(sheet_rows):
                if not any(cells) or cells[0].startswith('_ftnref'):continue
                records.append({'publication_year':publication_year,'source_file':name,
                  'sheet':sheet.title.strip(),'excel_row':first_row+offset,
                  'source_range':address,'values_json':json.dumps(cells,ensure_ascii=False),
                  'row_text':' | '.join(x for x in cells if x)})
                count+=1
            sheet_index.append({'publication_year':publication_year,'source_file':name,
                'sheet':sheet.title.strip(),'used_range':address,'nonempty_rows':count})
        print(f'Loaded {publication_year}: {len(wb.worksheets)} sheets')
    return records,sheet_index,figures

def row_with(sheet,label,start=0):
    """First row whose cell at or after column `start` equals `label`."""
    for r in sheet:
        if any(c.strip()==label for c in r[start:]):return r
    raise ValueError('Row not found: '+label)

def cell_after(row,label):
    """Non-empty cells following the label cell, as numbers where possible."""
    i=next(i for i,c in enumerate(row) if c.strip()==label)
    return [None if c.strip()=='-' else float(c) for c in row[i+1:] if c]

def build_chart_series(figures):
    """Every chart series is read from a named workbook sheet, never typed in."""
    series=[]
    def add(chart,category,value,unit,sheet):
        series.append({'chart':chart,'category':category,'value':value,'unit':unit,
            'provenance':f'2025 workbook, {sheet}','verification_status':'read_from_workbook'})
    fig12=figures['Figure 12']
    for label,shown in [('<15','<15'),('15-19','15–19'),('20-24','20–24'),('25-29','25–29'),('30-34','30–34'),('35-39','35–39'),('40-44','40–44'),('45-49','45–49'),('50+','50+')]:
        add('maternal_age',shown,round(cell_after(row_with(fig12,label),label)[1],1),'percent','Figure 12')
    for label in ('15-19','20-24','25-29','30-34','35-39','40-44','45-49'):
        add('age_specific_fertility',label.replace('-','–'),round(cell_after(row_with(figures['Figure15'],label),label)[0],1),'births per 1000 women','Figure15')
    t27=figures['Table27']
    for label in ('0-4','5-9','10-14','15-19','20-24','25-29','30-34','35-39','40-44','45-49','50+'):
        add('divorce_duration',label.replace('-','–'),int(cell_after(row_with(t27,label),label)[2]),'divorces','Table27')
    t24=figures['Table24']
    head=next(r for r in t24 if 'Eastern Province' in [c.strip() for c in r])
    columns={c.strip():i for i,c in enumerate(head)}
    places=[('Kigali City','Kigali City'),('Eastern Province','Eastern'),('Southern Province','Southern'),('Western Province','Western'),('Northern Province','Northern'),('Abroad','Abroad')]
    # The header row repeats these labels, so match on the row's first filled cell.
    t24_rows={next((c.strip() for c in r if c.strip()),''):r for r in t24 if r is not head}
    totals=t24_rows['Grand Total']
    for label,shown in places:
        wife=t24_rows[label]
        add('divorce_province_wife',shown,int(float(wife[columns['Grand Total']])),'divorces','Table24')
        add('divorce_province_husband',shown,int(float(totals[columns[label]])),'divorces','Table24')
    fig25=figures['Figure25']
    for label,shown in [('Noncommunicable diseases','Non-communicable'),('Communicable, maternal, perinatal and nutritional conditions','Communicable etc.'),('Injuries','Injuries')]:
        row=row_with(fig25,label)
        values=cell_after(row,label)
        add('cause_facility',shown,round(values[1]*100,1),'percent','Figure25')
        add('cause_community',shown,round(values[3]*100,1),'percent','Figure25')
    return series

def district_extremes(table5):
    """Highest and lowest district by births; excludes the total and Abroad rows."""
    start=next(i for i,r in enumerate(table5) if r[0].strip()=='Grand Total')
    districts=[(r[0].strip(),int(float(r[3]))) for r in table5[start+1:]
               if r[0].strip() and r[0].strip()!='Abroad' and not r[0].startswith('Source')]
    return max(districts,key=lambda d:d[1]),min(districts,key=lambda d:d[1])

def display_row(display,value,source):
    """`display` is the text the page shows; fail if the workbook value disagrees."""
    suffix='%' if display.endswith('%') else ''
    body=display.rstrip('%')
    if ',' in body:
        shown=f'{int(value):,}'
    else:
        shown=f'{value:.{len(body.split(".")[1]) if "." in body else 0}f}'
    if shown!=body:
        raise ValueError(f'Dashboard shows {display} but source gives {shown}: {source}')
    return {'display':display,'value':body.replace(',',''),'suffix':suffix,'source':source}

def extract_series(figure,label):
    """Get series by matched original label, NOT by assumed hard-coded values."""
    years=next((row for row in figure if sum(x in [str(y) for y in range(2019,2026)] for x in row)>=6),None)
    if years is None:raise ValueError('Cannot find year columns')
    index={i:int(x) for i,x in enumerate(years) if x in [str(y) for y in range(2019,2026)]}
    row=next((r for r in figure if r and label.lower() in r[0].lower()),None)
    if row is None:raise ValueError(f'Not found: {label}')
    return {y:float(row[i]) for i,y in index.items() if i<len(row) and row[i]}

def main():
    records,index,figures=read_workbooks()
    write_csv(OUT/'nisr_workbook_rows.csv',list(records[0]),records)
    write_csv(OUT/'sheet_inventory.csv',list(index[0]),index)
    fig7=figures['Figure7'];fig6=figures['Figure6']
    birth=extract_series(fig7,'Birth registration completeness')
    death=extract_series(fig7,'Death registration completeness')
    timely_birth=extract_series(fig6,'birth registered timely')
    timely_death=extract_series(fig6,'death registered timely')
    trends=[]
    for year in sorted(set(birth)&set(death)&set(timely_birth)&set(timely_death)):
        trends.append({'year':year,'birth_completeness_pct':birth[year],
             'death_completeness_pct':death[year],
             'completeness_gap_pp':round(birth[year]-death[year],3),
             'birth_timely_pct':round(timely_birth[year],3),
             'death_timely_pct':round(timely_death[year],3),
             'source_file':SOURCE_FILES[2025], 'completeness_sheet':'Figure7',
             'timeliness_sheet':'Figure6'})
    write_csv(OUT/'registration_trends_2019_2025.csv',list(trends[0]),trends)
    summary=figures['10.3']
    header=next(r for r in summary if r and r[0]=='Indicator')
    yearcols={i:int(v) for i,v in enumerate(header) if v.isdigit() and 2019<=int(v)<=2025}
    def get_summary(name):
        row=next((r for r in summary if r and r[0].strip()==name),None)
        if row is None:raise ValueError('Missing 10.3 indicator: '+name)
        return {year:float(row[i]) for i,year in yearcols.items()}
    registered=get_summary('Registered live births (number)')
    adjusted_rate=get_summary('Adjusted Crude birth rate (per 1,000 population)')
    birth_trend=[]
    previous=None
    for year in sorted(yearcols.values()):
        count=int(registered[year]);rate=adjusted_rate[year]
        birth_trend.append({'year':year,'registered_live_births':count,
            'adjusted_crude_birth_rate_per_1000':rate,
            'registered_births_yoy_change': '' if previous is None else count-previous,
            'registered_births_yoy_pct':'' if previous is None else round((count/previous-1)*100,2),
            'source_file':SOURCE_FILES[2025],'source_sheet':'10.3',
            'count_indicator':'Registered live births (number)',
            'rate_indicator':'Adjusted Crude birth rate (per 1,000 population)'})
        previous=count
    assert birth_trend[0]['registered_live_births']==313398
    assert birth_trend[-1]['registered_live_births']==356838
    assert birth_trend[0]['adjusted_crude_birth_rate_per_1000']==29.1
    assert birth_trend[-1]['adjusted_crude_birth_rate_per_1000']==27.2
    write_csv(OUT/'birth_trends_2019_2025.csv',list(birth_trend[0]),birth_trend)
    mortality=figures['10.5']
    mh=next(r for r in mortality if r and r[0]=='Indicator')
    myears={i:int(v) for i,v in enumerate(mh) if v.isdigit() and 2019<=int(v)<=2025}
    def mortality_series(label):
        row=next((r for r in mortality if r and r[0].strip()==label),None)
        if row is None:raise ValueError('Missing mortality measure: '+label)
        return {year:float(row[i]) for i,year in myears.items()}
    death_counts=mortality_series('Registered deaths (number)')
    death_rates=mortality_series('Crude death rate per 1,000 (Adjusted)')
    death_rows=[];prev=None
    for year in sorted(myears.values()):
        count=int(death_counts[year])
        death_rows.append({'year':year,'registered_deaths':count,'adjusted_crude_death_rate_per_1000':death_rates[year],
          'registered_deaths_yoy_change':'' if prev is None else count-prev,
          'registered_deaths_yoy_pct':'' if prev is None else round((count/prev-1)*100,2),
          'source_file':SOURCE_FILES[2025],'source_sheet':'10.5'})
        prev=count
    assert death_rows[0]['registered_deaths']==23791 and death_rows[-1]['registered_deaths']==39355
    write_csv(OUT/'death_trends_2019_2025.csv',list(death_rows[0]),death_rows)
    marriage_sheet=figures['Table18']
    marriage_rows=[];prev=None
    for r in marriage_sheet:
        if not r or not r[0].isdigit() or not (2019<=int(r[0])<=2025):continue
        year=int(r[0]);count=int(float(r[1]));rate=float(r[3])
        marriage_rows.append({'year':year,'registered_marriages':count,'crude_marriage_rate_per_1000':rate,
          'registered_marriages_yoy_change':'' if prev is None else count-prev,
          'registered_marriages_yoy_pct':'' if prev is None else round((count/prev-1)*100,2),
          'source_file':SOURCE_FILES[2025],'source_sheet':'Table18',
          'source_discrepancy':'2025 factsheet has 50,236' if year==2025 else ''})
        prev=count
    assert len(marriage_rows)==7 and marriage_rows[-1]['registered_marriages']==50256
    write_csv(OUT/'marriage_trends_2019_2025.csv',list(marriage_rows[0]),marriage_rows)


    table9=figures['Table9']
    fig24=row_with(figures['Figure24'],'Usable')
    fig12=figures['Figure 12']
    adolescent=sum(cell_after(row_with(fig12,x),x)[1] for x in ('<15','15-19'))
    facility_deaths=cell_after(row_with(figures['Table 15'],'2025'),'2025')[0]
    registered_deaths=death_rows[-1]['registered_deaths']
    catalog=[
      ('Registered live births',birth_trend[-1]['registered_live_births'],'events','2025 workbook','10.3','Births occurring in 2025 registered in CRVS'),
      ('Registered deaths',registered_deaths,'events','2025 workbook','10.5','Deaths occurring and registered in 2025'),
      ('Registered marriages',marriage_rows[-1]['registered_marriages'],'events','2025 workbook','Table18; the 2025 factsheet prints 50,236','Marriages registered in CRVS'),
      ('Registered divorces',int(cell_after(row_with(figures['Figure32'],'Grand Total'),'Grand Total')[0]),'events','2025 workbook','Figure32','Divorces recorded in CRVS; not a divorce rate'),
      ('Divorces granted and registered in 2025',int(cell_after(row_with(figures['Table28'],'Number of divorces granted and registered in 2025'),'Number of divorces granted and registered in 2025')[0]),'events','2025 workbook','Table28','Subset of legally granted and registered divorces'),
      ('Live births in health facilities',cell_after(row_with(figures['Table4'],'All'),'All')[4],'percent','2025 workbook','Table4','Of registered live births'),
      ('Births to adolescent mothers aged 10–19',round(adolescent,1),'percent','2025 workbook','Figure 12, <15 plus 15-19','Of registered live births'),
      ('Total fertility rate',cell_after(row_with(table9,'TFR'),'TFR')[6],'births per woman','2025 workbook','Table9','Period total fertility rate'),
      ('Registered community deaths',round(100*(registered_deaths-facility_deaths)/registered_deaths,1),'percent','2025 workbook','10.5 and Table 15, deaths not registered at facilities','Of registered deaths'),
      ('Usable cause of death in facilities',float(fig24[8]),'percent','2025 workbook','Figure24; Figure 23 swaps the community labels','Of relevant certified facility deaths'),
      ('Usable cause of death in community',float(fig24[11]),'percent','2025 workbook','Figure24 and 10.9; Figure 23 swaps the community labels','Of relevant community causes'),
      ('Community of property marriages',round(cell_after(row_with(figures['Table20'],'Community of property regime'),'Community of property regime')[1],1),'percent','2025 workbook','Table20','Of registered marriages'),
    ]
    indicators=[{'year':2025,'indicator':a,'value':b,'unit':c,'source_document':d,'source_location':e,'definition':f} for a,b,c,d,e,f in catalog]
    write_csv(OUT/'factsheet_indicators_2025.csv',list(indicators[0]),indicators)
    a=trends[-1]; prev=trends[-2]; base=trends[0]
    findings=[
      {'finding':'Birth registration completeness, 2025','value':a['birth_completeness_pct'],'unit':'percent','calculation':'2025 source series','source':'2025 workbook, Figure7'},
      {'finding':'Death registration completeness, 2025','value':a['death_completeness_pct'],'unit':'percent','calculation':'2025 source series','source':'2025 workbook, Figure7'},
      {'finding':'Completeness gap, 2025','value':a['completeness_gap_pp'],'unit':'percentage points','calculation':'birth completeness minus death completeness','source':'2025 workbook, Figure7'},
      {'finding':'Birth completeness change since 2019','value':round(a['birth_completeness_pct']-base['birth_completeness_pct'],2),'unit':'percentage points','calculation':'2025 minus 2019','source':'2025 workbook, Figure7'},
      {'finding':'Death completeness change since 2019','value':round(a['death_completeness_pct']-base['death_completeness_pct'],2),'unit':'percentage points','calculation':'2025 minus 2019','source':'2025 workbook, Figure7'},
      {'finding':'Birth completeness change from 2024','value':round(a['birth_completeness_pct']-prev['birth_completeness_pct'],2),'unit':'percentage points','calculation':'2025 minus 2024','source':'2025 workbook, Figure7'},
      {'finding':'Death completeness change from 2024','value':round(a['death_completeness_pct']-prev['death_completeness_pct'],2),'unit':'percentage points','calculation':'2025 minus 2024','source':'2025 workbook, Figure7'},
    ]
    write_csv(OUT/'dashboard_findings.csv',list(findings[0]),findings)
    series=build_chart_series(figures)
    write_csv(OUT/'dashboard_supplemental_source.csv',list(series[0]),series)
    def chart_value(chart,category):
        return next(s['value'] for s in series if s['chart']==chart and s['category']==category)
    top,bottom=district_extremes(figures['Table5'])
    assert (top[0],bottom[0])==('Gasabo','Nyaruguru')
    certified=int(cell_after(row_with(figures['Table 15'],'2025'),'2025')[1])
    verbal_autopsy=int(cell_after(row_with(figures['Table 16'],'VA for deaths that occurred in the reporting year'),'VA for deaths that occurred in the reporting year')[3])
    by_name={i['indicator']:i['value'] for i in indicators}
    fact='2025 workbook, '
    display=[
      display_row('99.1%',a['birth_timely_pct'],fact+'Figure6 timely registered births'),
      display_row('89.9%',a['death_timely_pct'],fact+'Figure6 timely registered deaths'),
      display_row('8.1%',cell_after(row_with(table9,'Low birth weight (%)'),'Low birth weight (%)')[6],fact+'Table9 low birth weight'),
      display_row('24,393',top[1],fact+'Table5, highest district by residence of mother'),
      display_row('8,229',bottom[1],fact+'Table5, lowest district by residence of mother'),
      display_row('102.7',cell_after(row_with(table9,'Sex ratio at Birth'),'Sex ratio at Birth')[6],fact+'Table9 sex ratio at birth'),
      display_row('14,997',certified,fact+'Table 15 medically certified deaths'),
      display_row('18,921',verbal_autopsy,fact+'Table 16 verbal autopsies, deaths in the reporting year'),
      display_row('87.0%',base['birth_completeness_pct'],fact+'Figure7, 2019 birth completeness'),
      display_row('92.9%',a['birth_completeness_pct'],fact+'Figure7, 2025 birth completeness'),
      display_row('50.5%',a['death_completeness_pct'],fact+'Figure7, 2025 death completeness'),
      display_row('31.4%',base['death_completeness_pct'],fact+'Figure7, 2019 death completeness'),
      display_row('42.4',a['completeness_gap_pp'],fact+'Figure7, 2025 birth minus death completeness'),
      display_row('19.1',a['death_completeness_pct']-base['death_completeness_pct'],fact+'Figure7, death completeness 2025 minus 2019'),
      display_row('4.4',a['death_completeness_pct']-prev['death_completeness_pct'],fact+'Figure7, death completeness 2025 minus 2024'),
      display_row('5.9',a['birth_completeness_pct']-base['birth_completeness_pct'],fact+'Figure7, birth completeness 2025 minus 2019; also the 2019 crude death rate in 10.5'),
      display_row('5.5',death_rows[-1]['adjusted_crude_death_rate_per_1000'],fact+'10.5, 2025 adjusted crude death rate'),
      display_row('29.1',birth_trend[0]['adjusted_crude_birth_rate_per_1000'],fact+'10.3, 2019 adjusted crude birth rate'),
      display_row('27.2',birth_trend[-1]['adjusted_crude_birth_rate_per_1000'],fact+'10.3, 2025 adjusted crude birth rate'),
      display_row('1.9',birth_trend[0]['adjusted_crude_birth_rate_per_1000']-birth_trend[-1]['adjusted_crude_birth_rate_per_1000'],fact+'10.3, adjusted crude birth rate 2019 minus 2025'),
      display_row('3.9',marriage_rows[0]['crude_marriage_rate_per_1000'],fact+'Table18, 2019 crude marriage rate'),
      display_row('3.6',marriage_rows[-1]['crude_marriage_rate_per_1000'],fact+'Table18, 2025 crude marriage rate'),
      display_row('23,791',death_rows[0]['registered_deaths'],fact+'10.5, 2019 registered deaths'),
      display_row('43,440',birth_trend[-1]['registered_live_births']-birth_trend[0]['registered_live_births'],fact+'10.3, registered births 2025 minus 2019'),
      display_row('13.9%',100*(birth_trend[-1]['registered_live_births']/birth_trend[0]['registered_live_births']-1),fact+'10.3, registered births change 2019 to 2025'),
      display_row('166.3',chart_value('age_specific_fertility','25–29'),fact+'Figure15, age 25-29'),
      display_row('49.5%',chart_value('cause_facility','Non-communicable'),fact+'Figure25, facility non-communicable share'),
      display_row('58.8%',chart_value('cause_community','Non-communicable'),fact+'Figure25, community non-communicable share'),
      display_row('24%',100*chart_value('divorce_duration','5–9')/by_name['Registered divorces'],fact+'Table27, 5-9 year share of divorces'),
      display_row('356,838',by_name['Registered live births'],fact+'10.3 registered live births'),
      display_row('39,355',by_name['Registered deaths'],fact+'10.5 registered deaths'),
      display_row('50,256',by_name['Registered marriages'],fact+'Table18 registered marriages'),
      display_row('4,479',by_name['Registered divorces'],fact+'Figure32 registered divorces'),
      display_row('2,629',by_name['Divorces granted and registered in 2025'],fact+'Table28 divorces granted and registered'),
      display_row('97.8%',by_name['Live births in health facilities'],fact+'Table4 facility births'),
      display_row('8.5%',by_name['Births to adolescent mothers aged 10–19'],fact+'Figure 12, mothers under 20'),
      display_row('3.5',by_name['Total fertility rate'],fact+'Table9 total fertility rate'),
      display_row('55.9%',by_name['Registered community deaths'],fact+'10.5 and Table 15, community share of registered deaths'),
      display_row('82.3%',by_name['Usable cause of death in facilities'],fact+'Figure24 usable facility causes'),
      display_row('83.9%',by_name['Usable cause of death in community'],fact+'Figure24 usable community causes'),
      display_row('97.0%',by_name['Community of property marriages'],fact+'Table20 community of property'),
    ]
    assert death_rows[0]['adjusted_crude_death_rate_per_1000']==5.9
    write_csv(OUT/'dashboard_display_source.csv',list(display[0]),display)
    assert len(trends)==7 and a['birth_completeness_pct']==92.9 and a['death_completeness_pct']==50.5
    assert a['completeness_gap_pp']==42.4
    report={'raw_rows':len(records),'source_sheets':len(index),'time_series_years':len(trends),
            'validation':'PASS','checks':['7 calendar years present','2025 births 92.9%','2025 deaths 50.5%','2025 gap 42.4 percentage points'],
            'limitations':['Some source sheets include formatted/empty ranges and mixed units; raw rows must be interpreted with sheet context',
               'The 2025 marriage factsheet figure (50,236) is quoted from the published factsheet, which is not stored in this repository',
               'Cross-year comparability depends on NISR definitions and revisions; latest 2025 workbook is used for historical trends', '2025 marriage count discrepancy: 50,256 in workbook Table18 versus 50,236 in 2025 marriage factsheet; chart preserves Table18']}
    (OUT/'validation_report.json').write_text(json.dumps(report,indent=2),encoding='utf-8')
    try:
        import matplotlib
        matplotlib.use('Agg')
        import matplotlib.pyplot as plt
        fig,ax=plt.subplots(figsize=(9,4.8))
        ax.plot([r['year'] for r in trends],[r['birth_completeness_pct'] for r in trends],marker='o',label='Birth registration')
        ax.plot([r['year'] for r in trends],[r['death_completeness_pct'] for r in trends],marker='o',label='Death registration')
        ax.set(ylabel='Registration completeness (%)',xlabel='Year',ylim=(0,100),title='Rwanda: registration completeness, 2019–2025')
        ax.legend(frameon=False);ax.spines[['top','right']].set_visible(False);fig.tight_layout()
        fig.savefig(OUT/'completeness_trend.png',dpi=180);plt.close(fig)
    except ImportError:
        print('matplotlib not installed; chart skipped')
    print(json.dumps(report,indent=2))

if __name__=='__main__':main()
